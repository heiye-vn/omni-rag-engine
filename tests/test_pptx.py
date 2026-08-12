import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from app.parsers.pptx import PPTXParser


def test_pptx_parser():
    print(">>> 1. 创建临时测试用 PPTX 文件 (sample.pptx)...")
    test_pptx_path = Path("tests/sample.pptx")
    test_pptx_path.parent.mkdir(exist_ok=True)

    prs = Presentation()

    # --- Slide 1: 标题页 ---
    slide_layout = prs.slide_layouts[0]  # 标题布局
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.placeholders[0].text = "项目汇报"
    slide1.placeholders[1].text = "2026 年度技术架构升级方案"

    # --- Slide 2: 内容页 (包含文本框 + 表格) ---
    blank_layout = prs.slide_layouts[6]  # 空白布局
    slide2 = prs.slides.add_slide(blank_layout)

    # 添加文本框
    from pptx.util import Inches, Pt
    txBox = slide2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "核心技术指标"
    p.font.bold = True
    p.font.size = Pt(24)

    p2 = tf.add_paragraph()
    p2.text = "以下为各模块性能对比数据"

    # 添加表格
    rows, cols = 3, 3
    table_shape = slide2.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(8), Inches(2))
    table = table_shape.table
    table_data = [
        ["模块", "延迟(ms)", "吞吐量(QPS)"],
        ["认证服务", "12", "5000"],
        ["数据服务", "45", "2000"],
    ]
    for r_idx, row_data in enumerate(table_data):
        for c_idx, cell_text in enumerate(row_data):
            table.cell(r_idx, c_idx).text = cell_text

    prs.save(test_pptx_path)

    print(">>> 2. 使用 PPTXParser 进行解析...")
    parser = PPTXParser()
    parsed_doc = parser.parse(str(test_pptx_path))

    print(f"文件名: {parsed_doc.file_name}")
    print(f"文件类型: {parsed_doc.file_type}")
    print(f"总页数: {parsed_doc.total_pages}")
    print(f"解析到 Element 节点数: {len(parsed_doc.elements)}")

    assert parsed_doc.file_type == "pptx"
    assert parsed_doc.total_pages == 2

    # 验证各 Element 的页码与类型
    for i, elem in enumerate(parsed_doc.elements, 1):
        print(f"\n[{i}] type={elem.type}, page={elem.location.page_number if elem.location else 'N/A'}")
        content_preview = elem.content[:60] + "..." if len(elem.content) > 60 else elem.content
        print(f"    content={content_preview!r}")
        if elem.location and elem.location.bbox:
            print(f"    bbox={elem.location.bbox}")

    # 验证 Slide 1 的标题占位符被识别为 heading
    slide1_headings = [e for e in parsed_doc.elements if e.location and e.location.page_number == 1 and e.type == "heading"]
    assert len(slide1_headings) >= 1, "Slide 1 应至少包含一个标题元素"

    # 验证 Slide 2 的表格被提取
    tables = [e for e in parsed_doc.elements if e.type == "table"]
    assert len(tables) >= 1, "应至少包含一个表格元素"
    assert "<table" in tables[0].raw_content, "raw_content 应为 HTML 表格"

    # 验证 JSON 序列化
    doc_dict = parsed_doc.to_dict()
    doc_json = json.dumps(doc_dict, ensure_ascii=False, indent=2)
    assert "page_number" in doc_json
    print("\nJSON 序列化验证成功!")

    # 清理测试文件
    if test_pptx_path.exists():
        test_pptx_path.unlink()

    print("\n[OK] PPTXParser 测试全部成功通过!")


if __name__ == "__main__":
    test_pptx_parser()
