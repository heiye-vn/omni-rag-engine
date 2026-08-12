from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

from app.models import Element, Location, ParsedDocument
from .base import BaseParser


class XLSXParser(BaseParser):
    """XLSX / Excel 表格文档解析器（支持多工作表提取、Markdown/HTML 转换与单元格定位）"""

    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        # load_workbook data_only=True 用于读取计算后的单元格真实数值而非原始公式
        wb = openpyxl.load_workbook(file_path, data_only=True)

        elements: list[Element] = []
        raw_text_parts: list[str] = []
        sheet_names: list[str] = wb.sheetnames

        for sheet_name in sheet_names:
            sheet = wb[sheet_name]

            # 1. 提取非空单元格二维数据
            rows_data: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                # 将 None 转化为空字符串，其它数值转化为 str
                row_cells = [str(cell).strip() if cell is not None else "" for cell in row]
                if any(row_cells):  # 过滤全空行
                    rows_data.append(row_cells)

            if not rows_data:
                continue

            # 2. 裁切右侧多余的全空列
            max_col = max(
                (idx for row in rows_data for idx, val in enumerate(row) if val),
                default=0,
            ) + 1
            trimmed_rows = [row[:max_col] for row in rows_data]

            headers = trimmed_rows[0]

            # 3. 转换文本表示与 HTML 渲染字符串
            markdown_text = self._rows_to_markdown(trimmed_rows)
            html_table = self._rows_to_html(trimmed_rows)

            raw_text_parts.append(f"## Sheet: {sheet_name}\n\n{markdown_text}")

            # 4. 计算工作表单元格区域定位表达式 (例如 Sheet1!A1:D15)
            max_row_num = len(trimmed_rows)
            max_col_letter = get_column_letter(max_col)
            cell_range = f"{sheet_name}!A1:{max_col_letter}{max_row_num}"

            elements.append(
                Element(
                    type="table",
                    content=markdown_text,
                    raw_content=html_table,  # 可供前端直接渲染的 HTML <table> 结构
                    location=Location(selector=cell_range),
                    metadata={
                        "sheet_name": sheet_name,
                        "headers": headers,
                        "rows_count": len(trimmed_rows),
                        "cols_count": max_col,
                    },
                )
            )

        wb.close()

        return ParsedDocument(
            file_name=path.name,
            file_type="xlsx",
            raw_text="\n\n".join(raw_text_parts),
            file_path=str(path.absolute()),
            metadata={"sheet_names": sheet_names},
            elements=elements,
        )

    @staticmethod
    def _rows_to_markdown(rows: list[list[str]]) -> str:
        """将二维列表转换为 Markdown 表格格式"""
        if not rows:
            return ""

        header = rows[0]
        result = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(["---"] * len(header)) + " |",
        ]

        for row in rows[1:]:
            # 补齐不足列
            padded_row = row + [""] * (len(header) - len(row))
            result.append("| " + " | ".join(padded_row) + " |")

        return "\n".join(result)

    @staticmethod
    def _rows_to_html(rows: list[list[str]]) -> str:
        """将二维列表转换为 HTML <table> 格式（可被前端直接渲染）"""
        if not rows:
            return "<table></table>"

        lines = ["<table border='1'>"]
        header = rows[0]
        lines.append(
            "  <thead>\n    <tr>"
            + "".join(f"<th>{cell}</th>" for cell in header)
            + "</tr>\n  </thead>"
        )
        lines.append("  <tbody>")

        for row in rows[1:]:
            padded_row = row + [""] * (len(header) - len(row))
            lines.append(
                "    <tr>"
                + "".join(f"<td>{cell}</td>" for cell in padded_row)
                + "</tr>"
            )

        lines.append("  </tbody>\n</table>")
        return "\n".join(lines)
