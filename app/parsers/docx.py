from pathlib import Path

from docx import Document

from app.models import Element, Location, ParsedDocument

from .base import BaseParser


class DOCXParser(BaseParser):
    """DOCX 文档解析器（支持元素定位与 JSON 友好序列化）"""

    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        doc = Document(file_path)

        elements: list[Element] = []
        raw_texts: list[str] = []

        # ==========================
        # Paragraph（段落）
        # ==========================
        for idx, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue

            raw_texts.append(text)

            if paragraph.style.name.startswith("Heading"):
                element_type = "heading"
            else:
                element_type = "paragraph"

            # 导出可渲染/可传输的 XML 片段或原文本作为 raw_content
            raw_xml = paragraph._element.xml if hasattr(paragraph, "_element") else text

            elements.append(
                Element(
                    type=element_type,
                    content=text,
                    raw_content=raw_xml,
                    location=Location(selector=f"body > p[{idx+1}]"),
                    metadata={"style": paragraph.style.name},
                )
            )

        # ==========================
        # Table（表格）
        # ==========================
        for t_idx, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])

            markdown = self.table_to_markdown(rows)
            raw_texts.append(markdown)

            elements.append(
                Element(
                    type="table",
                    content=markdown,
                    raw_content=markdown,  # Markdown 表格文本可以直接在前端渲染
                    location=Location(selector=f"body > table[{t_idx+1}]"),
                    metadata={"rows_count": len(rows)},
                )
            )

        return ParsedDocument(
            file_name=path.name,
            file_type="docx",
            raw_text="\n\n".join(raw_texts),
            file_path=str(path.absolute()),
            elements=elements,
        )

    @staticmethod
    def table_to_markdown(rows: list[list[str]]) -> str:
        """将表格二维列表转换为 Markdown 格式文本"""
        if not rows:
            return ""

        header = rows[0]
        result = []

        result.append("| " + " | ".join(header) + " |")
        result.append("| " + " | ".join(["---"] * len(header)) + " |")

        for row in rows[1:]:
            result.append("| " + " | ".join(row) + " |")

        return "\n".join(result)
